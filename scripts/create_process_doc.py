#!/usr/bin/env python3
"""Generate Word document for RAG Demo Process Flow."""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_process_document():
    doc = Document()

    # Title
    title = doc.add_heading('E-commerce RAG Chatbot - Demo Process Flow', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('Complete step-by-step guide for the RAG Masterclass demonstration')
    doc.add_paragraph()

    # =========================================================================
    # PHASE 1: PREREQUISITES
    # =========================================================================
    doc.add_heading('Phase 1: Prerequisites & Setup', level=1)

    doc.add_heading('1.1 AWS Account Setup', level=2)
    doc.add_paragraph('Before starting the demo, ensure the following are configured:')

    items = [
        'AWS Account with appropriate permissions',
        'IAM User with programmatic access (Access Key ID + Secret Access Key)',
        'AWS CLI installed and configured locally',
        'Python 3.9+ installed',
        'Docker Desktop installed'
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    # UI Callout
    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('AWS Console > IAM > Users > Create user with programmatic access')

    doc.add_heading('1.2 Environment Configuration', level=2)
    doc.add_paragraph('Create a .env file in the project root with:')

    code = doc.add_paragraph()
    code.add_run('''AWS_ACCESS_KEY_ID=your_access_key
AWS_SECRET_ACCESS_KEY=your_secret_key
AWS_DEFAULT_REGION=us-east-1
KNOWLEDGE_BASE_ID=your_kb_id
LLM_BASE_URL=your_llm_endpoint
LLM_API_KEY=your_api_key
LLM_MODEL=gpt-4o''').italic = True

    # =========================================================================
    # PHASE 2: S3 & KNOWLEDGE BASE
    # =========================================================================
    doc.add_heading('Phase 2: S3 Bucket & Knowledge Base Setup', level=1)

    doc.add_heading('2.1 Create S3 Bucket', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('AWS Console > S3 > Create Bucket')

    steps = [
        'Navigate to S3 in AWS Console',
        'Click "Create bucket"',
        'Bucket name: ecom-rag-bucket',
        'Region: us-east-1',
        'Keep default settings (Block all public access)',
        'Click "Create bucket"'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('2.2 Upload Data to S3', level=2)
    doc.add_paragraph('Run Notebook 01.5 or the upload script to push FAQ data to S3:')

    p = doc.add_paragraph()
    p.add_run('[RUN NOTEBOOK] ').bold = True
    p.add_run('notebooks/01.5_upload_to_s3.ipynb')

    doc.add_paragraph('This creates a timestamped folder like: knowledge-base-20260117_143521/')

    doc.add_heading('2.3 Create Bedrock Knowledge Base', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('AWS Console > Amazon Bedrock > Knowledge bases')

    steps = [
        'Navigate to Amazon Bedrock > Knowledge bases',
        'Click "Create knowledge base"',
        'Name: ecommerce-faq-kb',
        'IAM role: Create new role (let AWS create it)',
        'Click "Next"'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('2.4 Configure Data Source', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('Knowledge Base > Data source configuration')

    steps = [
        'Data source name: s3-ecommerce-data',
        'S3 URI: s3://ecom-rag-bucket/knowledge-base-{timestamp}/',
        'Click "Next"'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('2.5 Select Embeddings Model', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('Knowledge Base > Embeddings model selection')

    steps = [
        'Embeddings model: Titan Embeddings G1 - Text v2',
        'Vector database: Quick create new vector store (OpenSearch Serverless)',
        'Click "Next" then "Create knowledge base"'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('2.6 Sync/Ingest Data', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('Knowledge Base > Data source > Sync')

    steps = [
        'Wait for Knowledge Base creation to complete',
        'Go to the Data source section',
        'Click "Sync" to start ingestion',
        'Wait for sync to complete (Status: Available)',
        'Copy the Knowledge Base ID (e.g., XVIAYNVNB2)'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run('Important: ').bold = True
    p.add_run('Update KNOWLEDGE_BASE_ID in your .env file with the new ID')

    # =========================================================================
    # PHASE 3: NOTEBOOKS
    # =========================================================================
    doc.add_heading('Phase 3: Run Demo Notebooks', level=1)

    doc.add_heading('3.1 Notebook 00: Setup & Connections', level=2)
    p = doc.add_paragraph()
    p.add_run('[RUN NOTEBOOK] ').bold = True
    p.add_run('notebooks/00_setup_and_connections.ipynb')

    items = [
        'Validates AWS credentials',
        'Tests LLM API connection',
        'Tests Titan Embeddings',
        'Downloads FAQ dataset'
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('3.2 Notebook 01: Data Exploration & Chunking', level=2)
    p = doc.add_paragraph()
    p.add_run('[RUN NOTEBOOK] ').bold = True
    p.add_run('notebooks/01_data_exploration_chunking.ipynb')

    items = [
        'Explores the FAQ dataset',
        'Demonstrates different chunking strategies',
        'Prepares data for Knowledge Base'
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('3.3 Notebook 02: Embeddings & Indexing', level=2)
    p = doc.add_paragraph()
    p.add_run('[RUN NOTEBOOK] ').bold = True
    p.add_run('notebooks/02_embeddings_and_indexing.ipynb')

    items = [
        'Generates embeddings with Titan',
        'Visualizes embedding space',
        'Explains vector similarity'
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    doc.add_heading('3.4 Notebook 03: Retrieval & RAG', level=2)
    p = doc.add_paragraph()
    p.add_run('[RUN NOTEBOOK] ').bold = True
    p.add_run('notebooks/03_retrieval_and_rag.ipynb')

    items = [
        'Demonstrates Vector Search',
        'Demonstrates BM25 Search',
        'Demonstrates Hybrid Search',
        'Shows conversational RAG with memory',
        'Source attribution with citations'
    ]
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

    # =========================================================================
    # PHASE 4: STREAMLIT APP
    # =========================================================================
    doc.add_heading('Phase 4: Streamlit Application', level=1)

    doc.add_heading('4.1 Run Locally', level=2)
    doc.add_paragraph('Test the Streamlit app locally before deployment:')

    code = doc.add_paragraph()
    code.add_run('streamlit run app.py').italic = True

    doc.add_paragraph('Open http://localhost:8501 in browser')

    p = doc.add_paragraph()
    p.add_run('[SHOW DEMO] ').bold = True
    p.add_run('Demonstrate chatbot features: ask questions, show memory, show sources')

    doc.add_heading('4.2 Build Docker Image', level=2)
    doc.add_paragraph('Build the Docker image for deployment:')

    code = doc.add_paragraph()
    code.add_run('docker build -t ecommerce-rag-chatbot .').italic = True

    # =========================================================================
    # PHASE 5: AWS DEPLOYMENT
    # =========================================================================
    doc.add_heading('Phase 5: AWS Deployment (ECR + Fargate)', level=1)

    doc.add_heading('5.1 Create ECR Repository', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('AWS Console > ECR > Create Repository')

    steps = [
        'Navigate to ECR (Elastic Container Registry)',
        'Click "Create repository"',
        'Repository name: ecommerce-rag-chatbot',
        'Click "Create"',
        'Copy the repository URI'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('5.2 Push Image to ECR', level=2)
    doc.add_paragraph('Run these commands to push the Docker image:')

    code = doc.add_paragraph()
    code.add_run('''# Login to ECR
aws ecr get-login-password --region us-east-1 | docker login --username AWS --password-stdin <account-id>.dkr.ecr.us-east-1.amazonaws.com

# Tag and push
docker tag ecommerce-rag-chatbot:latest <account-id>.dkr.ecr.us-east-1.amazonaws.com/ecommerce-rag-chatbot:latest
docker push <account-id>.dkr.ecr.us-east-1.amazonaws.com/ecommerce-rag-chatbot:latest''').italic = True

    doc.add_heading('5.3 Create ECS Cluster', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('AWS Console > ECS > Clusters > Create Cluster')

    steps = [
        'Navigate to ECS > Clusters',
        'Click "Create Cluster"',
        'Cluster name: ecommerce-rag-cluster',
        'Infrastructure: AWS Fargate (serverless)',
        'Click "Create"'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('5.4 Create Task Definition', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('AWS Console > ECS > Task Definitions > Create')

    steps = [
        'Click "Create new Task Definition"',
        'Family name: ecommerce-rag-task',
        'Launch type: AWS Fargate',
        'CPU: 0.5 vCPU, Memory: 1 GB',
        'Container name: ecommerce-rag-app',
        'Image URI: <your-ecr-uri>:latest',
        'Container port: 8501',
        'Click "Create"'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('5.5 Create Service', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('AWS Console > ECS > Cluster > Services > Create')

    steps = [
        'Go to your cluster > Services tab',
        'Click "Create"',
        'Launch type: FARGATE',
        'Task definition: ecommerce-rag-task',
        'Service name: ecommerce-rag-service',
        'Desired tasks: 1',
        'VPC: Select default VPC',
        'Subnets: Select public subnets',
        'Security group: Create new, allow port 8501 inbound',
        'Public IP: ENABLED',
        'Click "Create"'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    doc.add_heading('5.6 Access the Application', level=2)

    p = doc.add_paragraph()
    p.add_run('[SHOW UI] ').bold = True
    p.add_run('ECS > Cluster > Service > Tasks > Public IP')

    steps = [
        'Wait for task to reach "Running" state',
        'Click on the running task',
        'Find the Public IP in Configuration',
        'Open http://<public-ip>:8501 in browser'
    ]
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f'{i}. {step}')

    p = doc.add_paragraph()
    p.add_run('[SHOW DEMO] ').bold = True
    p.add_run('Demonstrate the live deployed chatbot!')

    # =========================================================================
    # SUMMARY
    # =========================================================================
    doc.add_heading('Quick Reference: UI Checkpoints', level=1)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'

    # Header row
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Step'
    header_cells[1].text = 'AWS Service'
    header_cells[2].text = 'Action'

    # Make header bold
    for cell in header_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # Data rows
    data = [
        ('1', 'S3', 'Create bucket: ecom-rag-bucket'),
        ('2', 'Bedrock', 'Create Knowledge Base'),
        ('3', 'Bedrock', 'Configure S3 data source'),
        ('4', 'Bedrock', 'Select Titan Embeddings'),
        ('5', 'Bedrock', 'Sync/Ingest data'),
        ('6', 'ECR', 'Create repository'),
        ('7', 'ECS', 'Create cluster'),
        ('8', 'ECS', 'Create task definition'),
        ('9', 'ECS', 'Create service'),
        ('10', 'ECS', 'Get public IP & demo'),
    ]

    for step, service, action in data:
        row = table.add_row().cells
        row[0].text = step
        row[1].text = service
        row[2].text = action

    # Save document
    output_path = '/Users/apalakkode/Library/CloudStorage/OneDrive-PayPal/Agentic RAG/ecommerce-chatbot/docs/RAG_Demo_Process_Flow.docx'
    doc.save(output_path)
    print(f"Document saved to: {output_path}")
    return output_path

if __name__ == '__main__':
    create_process_document()
